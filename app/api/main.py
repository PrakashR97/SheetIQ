
import asyncio
import shutil
import tempfile
import uuid
import csv

from pathlib import Path
from openpyxl import load_workbook
from fastapi import FastAPI, File, HTTPException, UploadFile
from fastapi.responses import FileResponse

from app.agents.crewai.excel_tool import excel_cleaning_tool


app = FastAPI(
    title="AI Excel Data Cleaning API",
    version="1.0.0",
)

# Store generated files temporarily for download
OUTPUT_DIR = Path(tempfile.gettempdir()) / "ai_excel_cleaner"
OUTPUT_DIR.mkdir(parents=True, exist_ok=True)


@app.get("/health")
def health():
    return {"status": "healthy"}


@app.post("/clean")
async def clean_excel(file: UploadFile = File(...)):

    # Validate Excel file extension
    filename = Path(file.filename or "").name

    if Path(filename).suffix.lower() != ".xlsx":
        raise HTTPException(
            status_code=400,
            detail="Only .xlsx files are supported by this endpoint.",
        )

    # Create a unique folder for each request
    job_id = uuid.uuid4().hex
    job_dir = OUTPUT_DIR / job_id
    job_dir.mkdir(parents=True, exist_ok=True)

    input_path = job_dir / filename

    # Save uploaded Excel file
    try:
        with input_path.open("wb") as destination:
            shutil.copyfileobj(file.file, destination)
    finally:
        await file.close()

    # Run existing CrewAI tool and AutoGen analysis
    result = await asyncio.to_thread(
        excel_cleaning_tool.run,
        file_path=str(input_path),
    )

    # Cleaned Excel output
    cleaned_path = input_path.with_name(
        input_path.stem + "_cleaned.xlsx"
    )

    # Cleaning report output
    report_path = input_path.with_name(
        input_path.stem + "_cleaning_report.csv"
    )

    if not cleaned_path.exists() or not report_path.exists():
        raise HTTPException(
            status_code=500,
            detail=f"Excel cleaning failed: {result}",
        )

    # NEW: Generate cleaned CSV from cleaned Excel
    cleaned_csv_path = input_path.with_name(
        input_path.stem + "_cleaned.csv"
    )

    try:
        workbook = load_workbook(
            cleaned_path,
            read_only=True,
            data_only=True,
        )

        try:
            # CSV supports only one sheet.
            # Export the first sheet.
            sheet = workbook.worksheets[0]

            with cleaned_csv_path.open(
                "w",
                newline="",
                encoding="utf-8-sig",
            ) as csv_file:

                writer = csv.writer(csv_file)

                for row in sheet.iter_rows(values_only=True):
                    writer.writerow(row)

        finally:
            workbook.close()

    except Exception as error:
        raise HTTPException(
            status_code=500,
            detail=f"Cleaned CSV generation failed: {error}",
        )

    # Return all three download URLs
    return {
        "job_id": job_id,
        "message": "Excel cleaning completed successfully",
        "result": str(result),

        "cleaned_excel_url": f"/download/{job_id}/excel",

        "cleaned_csv_url": f"/download/{job_id}/cleaned-csv",

        "csv_report_url": f"/download/{job_id}/csv",
    }


@app.get("/download/{job_id}/{file_type}")
def download_file(job_id: str, file_type: str):

    # Validate job ID
    if len(job_id) != 32 or any(
        character not in "0123456789abcdef"
        for character in job_id
    ):
        raise HTTPException(
            status_code=400,
            detail="Invalid job ID",
        )

    job_dir = OUTPUT_DIR / job_id

    if not job_dir.is_dir():
        raise HTTPException(
            status_code=404,
            detail="Job not found",
        )

    # Download cleaned Excel
    if file_type == "excel":

        files = list(job_dir.glob("*_cleaned.xlsx"))

        media_type = (
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )

    # NEW: Download cleaned CSV
    elif file_type == "cleaned-csv":

        files = list(job_dir.glob("*_cleaned.csv"))

        media_type = "text/csv"

    # Download cleaning report CSV
    elif file_type == "csv":

        files = list(job_dir.glob("*_cleaning_report.csv"))

        media_type = "text/csv"

    else:
        raise HTTPException(
            status_code=400,
            detail="Invalid file type",
        )

    if not files:
        raise HTTPException(
            status_code=404,
            detail="File not found",
        )

    return FileResponse(
        path=files[0],
        filename=files[0].name,
        media_type=media_type,
    )
