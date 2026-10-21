import re

from datetime import date, datetime, time

from decimal import Decimal, InvalidOperation

from pathlib import Path

from typing import Any

from openpyxl.cell.cell import MergedCell

from openpyxl.formula import Tokenizer

from openpyxl.utils.datetime import from_excel

from app.config import DEFAULT_CONFIG, CleaningConfig

from app.services.cleaning_rules import null_replacement

from app.services.report import CleaningReport

from app.services.row_cleaner import is_fully_blank_row

MAX_DATE = datetime(2999, 1, 1)

def to_excel_date_format(python_format: str) -> str:

    """Convert supported Python date formats to Excel formats."""

    formats = {

        "%m-%d-%Y": "mm-dd-yyyy",

        "%d/%m/%Y": "dd/mm/yyyy",

        "%m/%d/%Y": "mm/dd/yyyy",

        "%Y-%m-%d": "yyyy-mm-dd",

    }

    if python_format not in formats:

        raise ValueError(

            f"Unsupported date format: {python_format}"

        )

    return formats[python_format]

INTEGER_HINTS = {

    "quantity", "qty", "count", "age", "units",

    "number of", "total count", "stock",

}

DECIMAL_HINTS = {

    "amount", "price", "salary", "cost", "rate",

    "balance", "revenue", "expense", "payment",

    "discount", "percentage", "percent", "total",

}

ID_HINTS = {

    "id", "identifier", "code", "zip", "postal",

    "phone", "mobile", "account", "reference",

}

def normalize_header(header: Any) -> str:

    text = str(header or "").strip().lower()

    text = re.sub(r"[_-]+", " ", text)

    return re.sub(r"\s+", " ", text)

def is_birth_date_column(header: Any) -> bool:

    name = normalize_header(header)

    return bool(re.search(

        r"\b(dob|birthdate|birthday|date of birth|"

        r"birth date|birth dt)\b",

        name,

    ))

def is_date_column(header: Any) -> bool:

    name = normalize_header(header)

    return (

        is_birth_date_column(header)

        or bool(re.search(

            r"\b(date|datetime|timestamp|dt)\b",

            name,

        ))

    )

def is_identifier_column(header: Any) -> bool:

    name = normalize_header(header)

    return any(

        re.search(rf"\b{re.escape(hint)}\b", name)

        for hint in ID_HINTS

    )

def is_missing(value: Any) -> bool:

    return value is None or (

        isinstance(value, str) and not value.strip()

    )

def parse_date_value(

    value: Any,

    epoch: datetime,

) -> datetime | None:

    """Parse supported Excel dates and common text dates."""

    if isinstance(value, datetime):

        return value

    if isinstance(value, date):

        return datetime.combine(value, time.min)

    if isinstance(value, bool):

        return None

    if isinstance(value, (int, float)):

        try:

            parsed = from_excel(value, epoch=epoch)

            return parsed if isinstance(parsed, datetime) else None

        except (ValueError, TypeError, OverflowError):

            return None

    if isinstance(value, str):

        text = value.strip()

        if not text:

            return None

        formats = (

            "%Y-%m-%d",

            "%Y/%m/%d",

            "%d-%m-%Y",

            "%d/%m/%Y",

            "%m-%d-%Y",

            "%m/%d/%Y",

            "%Y-%m-%d %H:%M:%S",

            "%d-%m-%Y %H:%M:%S",

            "%d/%m/%Y %H:%M:%S",

            "%m/%d/%Y %H:%M:%S",

        )

        for fmt in formats:

            try:

                return datetime.strptime(text, fmt)

            except ValueError:

                continue

    return None

def numeric_kind(value: Any) -> str | None:

    """Identify numeric values without converting identifiers."""

    if isinstance(value, bool) or is_missing(value):

        return None

    if isinstance(value, int):

        return "integer"

    if isinstance(value, (float, Decimal)):

        try:

            number = Decimal(str(value))

            return "decimal" if number % 1 else "integer"

        except InvalidOperation:

            return None

    if isinstance(value, str):

        text = value.strip().replace(",", "")

        if not re.fullmatch(r"[+-]?\d+(\.\d+)?", text):

            return None

        try:

            Decimal(text)

            return "decimal" if "." in text else "integer"

        except InvalidOperation:

            return None

    return None

def detect_column_type(

    header: Any,

    values: list[Any],

) -> str:

    """Detect date, integer, decimal, identifier, or text."""

    name = normalize_header(header)

    if is_date_column(header):

        return "date"

    if is_identifier_column(header):

        return "identifier"

    # Infer dates from actual values when the header gives no date hint.
    # Do not interpret plain numbers as Excel date serials here.
    nonblank_values = [value for value in values if not is_missing(value)]
    if nonblank_values:
        date_count = sum(
            1
            for value in nonblank_values
            if isinstance(value, (date, datetime))
            or (
                isinstance(value, str)
                and parse_date_value(value, datetime(1899, 12, 30)) is not None
            )
        )
        if date_count / len(nonblank_values) >= 0.8:
            return "date"

    if any(

        re.search(rf"\b{re.escape(hint)}\b", name)

        for hint in DECIMAL_HINTS

    ):

        return "decimal"

    if any(

        re.search(rf"\b{re.escape(hint)}\b", name)

        for hint in INTEGER_HINTS

    ):

        return "integer"

    nonblank = [

        value for value in values

        if not is_missing(value)

    ]

    kinds = [

        numeric_kind(value)

        for value in nonblank

        if not isinstance(value, (date, datetime, bool))

    ]

    kinds = [kind for kind in kinds if kind is not None]

    if kinds and len(kinds) >= len(nonblank) * 0.7:

        return "decimal" if "decimal" in kinds else "integer"

    return "text"

def _row_signature(

    values: tuple[Any, ...],

) -> tuple[Any, ...]:

    """Build a hashable signature for exact duplicate detection."""

    result = []

    for value in values:

        if isinstance(value, datetime):

            result.append(("datetime", value.isoformat()))

        elif isinstance(value, date):

            result.append(("date", value.isoformat()))

        elif isinstance(value, float):

            result.append(("float", value))

        elif isinstance(value, (str, int, bool, type(None))):

            result.append((type(value).__name__, value))

        else:

            result.append((type(value).__name__, repr(value)))

    return tuple(result)

# Deliberately support only plain, same-sheet A1 references initially.

# More complex Excel dependencies must not be silently rewritten.

_SIMPLE_REF = re.compile(r"(\$?[A-Za-z]{1,3})(\$?)([1-9][0-9]*)\Z")

def _rewrite_formulas_for_deleted_rows(workbook, worksheet, rows_to_delete):

    """Rewrite simple same-sheet A1 references before physical deletion.

    Raises ValueError for unsupported references or references to deleted rows.

    No workbook is saved if validation fails.

    """

    deleted = set(rows_to_delete)

    # Other sheets may depend on this one; don't risk breaking those links.

    for other in workbook.worksheets:

        if other is worksheet:

            continue

        for row in other.iter_rows():

            for cell in row:

                if cell.data_type == "f" and "!" in cell.value:

                    raise ValueError(

                        f"Unsupported cross-sheet formula at {other.title}!{cell.coordinate}"

                    )

    updates = []

    for row in worksheet.iter_rows():

        for cell in row:

            if cell.data_type != "f" or cell.row in deleted:

                continue

            try:

                tokens = Tokenizer(cell.value).items

            except Exception as exc:

                raise ValueError(f"Cannot parse formula at {worksheet.title}!{cell.coordinate}: {exc}") from exc

            parts = []

            for token in tokens:

                if token.type == "OPERAND" and token.subtype == "RANGE":

                    match = _SIMPLE_REF.fullmatch(token.value)

                    if not match:

                        raise ValueError(

                            f"Unsupported formula reference '{token.value}' "

                            f"at {worksheet.title}!{cell.coordinate}"

                        )

                    col, absolute, row_text = match.groups()

                    original_row = int(row_text)

                    if original_row in deleted:

                        raise ValueError(

                            f"Formula at {worksheet.title}!{cell.coordinate} "

                            f"references deleted row {original_row}"

                        )

                    new_row = original_row - sum(r < original_row for r in deleted)

                    parts.append(f"{col}{absolute}{new_row}")

                else:

                    parts.append(token.value)

            updates.append((cell, "=" + "".join(parts)))

    for cell, formula in updates:

        cell.value = formula

def clean_workbook(

    input_path: str | Path,

    output_path: str | Path | None = None,

    config: CleaningConfig = DEFAULT_CONFIG,

) -> CleaningReport:

    """Clean an Excel workbook without overwriting the original."""

    source = Path(input_path).resolve()

    if not source.is_file():

        raise FileNotFoundError(

            f"Input file not found: {source}"

        )

    if source.suffix.lower() not in {".xlsx", ".xlsm"}:

        raise ValueError(

            "Only .xlsx and .xlsm files are supported."

        )

    output = (

        Path(output_path).resolve()

        if output_path is not None

        else source.with_name(

            f"{source.stem}{config.output_suffix}{source.suffix}"

        )

    )

    if output == source:

        raise ValueError(

            "Output must not overwrite the original file."

        )

    if output.suffix.lower() != source.suffix.lower():

        raise ValueError(

            "Input and output extensions must match."

        )

    from app.services.excel_reader import read_workbook

    loaded = read_workbook(source)

    workbook = loaded["workbook"]

    report = CleaningReport(input_file=str(source))

    try:

        for worksheet in workbook.worksheets:

            report.sheets_processed += 1

            rows = list(worksheet.iter_rows())

            if not rows:

                continue

            # Ignore trailing formatting-only rows.

            last_data_row = 0

            for row_index, row in enumerate(rows, start=1):

                if not is_fully_blank_row(

                    cell.value for cell in row

                ):

                    last_data_row = row_index

            if last_data_row == 0:

                continue

            rows = rows[:last_data_row]

            headers = [cell.value for cell in rows[0]]

            # Detect column types before cleaning.

            column_types = []

            for column_index, header in enumerate(headers):

                values = [

                    row[column_index].value

                    for row in rows[1:]

                    if column_index < len(row)

                    and row[column_index].data_type != "f"

                ]

                column_types.append(

                    detect_column_type(header, values)

                )

            seen: set[tuple[Any, ...]] = set()

            rows_to_delete: list[int] = []

            for row_index, row in enumerate(

                rows[1:],

                start=2,

            ):

                values = tuple(

                    cell.value for cell in row

                )

                report.rows_processed += 1

                # Remove fully blank rows.

                if (

                    config.remove_blank_rows

                    and is_fully_blank_row(values)

                ):

                    rows_to_delete.append(row_index)

                    report.blank_rows_removed += 1

                    continue

                # Remove exact duplicates before cleaning.

                if config.remove_exact_duplicates:

                    signature = _row_signature(values)

                    if signature in seen:

                        rows_to_delete.append(row_index)

                        report.duplicate_rows_removed += 1

                        continue

                    seen.add(signature)

                for column_index, cell in enumerate(row):

                    if (

                        isinstance(cell, MergedCell)

                        or cell.data_type == "f"

                    ):

                        continue

                    value = cell.value

                    header = headers[column_index]

                    column_type = column_types[column_index]

                    replacement = value

                    reason = None

                    # DATE COLUMNS

                    if column_type == "date":

                        parsed = parse_date_value(

                            value,

                            workbook.epoch,

                        )

                        if parsed is None:

                            if is_birth_date_column(header):

                                replacement = "N/A"

                                cell.number_format = "General"

                            else:

                                replacement = MAX_DATE

                                #cell.number_format = config.date_format

                                cell.number_format = to_excel_date_format(config.date_format)

                            reason = "Missing or invalid date replaced"

                        else:

                            replacement = parsed

                            #cell.number_format = config.date_format

                            cell.number_format = to_excel_date_format(config.date_format)

                            if replacement != value:

                                reason = "Date normalized"

                    # INTEGER COLUMNS

                    elif column_type == "integer":

                        if is_missing(value):

                            replacement = 0

                            reason = "Missing integer replaced"

                        elif numeric_kind(value) is None:

                            report.invalid_values.append({

                                "sheet": worksheet.title,

                                "row": row_index,

                                "column": header,

                                "value": str(value),

                                "reason": "Suspicious integer value",

                            })

                    # DECIMAL COLUMNS

                    elif column_type == "decimal":

                        cell.number_format = "0.00"

                        if is_missing(value):

                            replacement = 0.0

                            reason = "Missing decimal replaced"

                        elif numeric_kind(value) is None:

                            report.invalid_values.append({

                                "sheet": worksheet.title,

                                "row": row_index,

                                "column": header,

                                "value": str(value),

                                "reason": "Suspicious decimal value",

                            })

                    # TEXT AND IDENTIFIER COLUMNS

                    else:

                        if is_missing(value):

                            replacement = null_replacement(

                                header,

                                value,

                            )

                            reason = "Missing text replaced"

                        elif (

                            isinstance(value, str)

                            and config.trim_whitespace

                        ):

                            replacement = value.strip()

                            if replacement != value:

                                reason = "Trimmed whitespace"

                                report.values_trimmed += 1

                    if replacement != value:

                        cell.value = replacement

                        report.changes.append({

                            "sheet": worksheet.title,

                            "row": row_index,

                            "column": header,

                            "original_value": str(value),

                            "new_value": str(replacement),

                            "reason": reason or "Value cleaned",

                        })

                        if is_missing(value):

                            report.null_values_replaced += 1

            if rows_to_delete:

                _rewrite_formulas_for_deleted_rows(

                    workbook, worksheet, rows_to_delete

                )

            # Formulas already reference their final target rows.

            for row_index in reversed(rows_to_delete):

                worksheet.delete_rows(row_index)

        output.parent.mkdir(

            parents=True,

            exist_ok=True,

        )

        workbook.save(output)

        report.output_file = str(output)

    except Exception as exc:

        report.errors.append(str(exc))

        raise

    finally:

        workbook.close()

    return report

