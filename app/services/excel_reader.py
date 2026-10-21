
from pathlib import Path

from openpyxl import load_workbook


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}


def read_workbook(file_path: str | Path) -> dict:
    """Read all worksheets while retaining workbook cell values and metadata."""
    path = Path(file_path)

    if not path.is_file():
        raise FileNotFoundError(f"Excel file not found: {path}")

    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError(
            "Supported formats: .xlsx and .xlsm. "
            "Legacy .xls files are not supported."
        )

    workbook = load_workbook(
        filename=path,
        read_only=False,
        data_only=False,
        keep_vba=path.suffix.lower() == ".xlsm",
    )

    return {
        "path": path,
        "workbook": workbook,
        "sheet_names": workbook.sheetnames.copy(),
    }

