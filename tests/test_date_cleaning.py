from datetime import datetime
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from openpyxl import Workbook, load_workbook

from app.services.excel_cleaner import (
    detect_column_type,
    clean_workbook,
)


def test_date_auto_detection():
    values = [
        "2026-10-21",
        "2026-10-22",
        "2026-10-23",
    ]

    assert detect_column_type("Field D", values) == "date"


def test_mixed_date_detection():
    values = [
        "2026-10-21",
        "2026-10-22",
        None,
        "2026-10-23",
        "invalid",
        "2026-10-24",
    ]

    assert detect_column_type("Field D", values) == "date"


def test_invalid_date_replacement():
    with TemporaryDirectory() as folder:
        source = Path(folder) / "dates.xlsx"

        wb = Workbook()
        ws = wb.active

        ws.append(["Field D"])
        ws.append(["2026-10-21"])
        ws.append(["2026-10-22"])
        ws.append(["2026-10-23"])
        ws.append(["2026-10-24"])
        ws.append(["invalid"])

        wb.save(source)
        wb.close()

        report = clean_workbook(source)

        result = load_workbook(report.output_file)

        try:
            assert result.active["A6"].value == datetime(2999, 1, 1)
        finally:
            result.close()


@pytest.mark.parametrize(
    "header, values, expected",
    [
        ("Field A", [10, 20, 30], "integer"),
        ("Field B", [10.5, 20.75], "decimal"),
        ("Field C", ["Alice", "Bob"], "text"),
        ("Customer ID", ["00123", "00456"], "identifier"),
        ("Order Date", ["2026-10-21"], "date"),
    ],
)
def test_column_type_detection(header, values, expected):
    assert detect_column_type(header, values) == expected