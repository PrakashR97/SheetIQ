from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from app.services.excel_cleaner import clean_workbook


def test_excel_cleaner():
    with TemporaryDirectory() as temp_dir:
        folder = Path(temp_dir)
        source = folder / "test_data.xlsx"
        output = folder / "sample_cleaned.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Customers"

        sheet.append(["Customer Name", "Amount"])
        sheet.append([" Alice ", 100])
        sheet.append([" Alice ", 100])  # Exact duplicate
        sheet.append([None, None])      # Blank row
        sheet.append(["Bob", None])

        second_sheet = workbook.create_sheet("Orders")
        second_sheet.append(["Order ID", "Status"])
        second_sheet.append([101, " Completed "])

        workbook.save(source)
        workbook.close()

        report = clean_workbook(source, output)

        assert output.exists()
        assert source.exists()
        assert report.sheets_processed == 2
        assert report.blank_rows_removed == 1
        assert report.duplicate_rows_removed == 1
        assert report.null_values_replaced == 1
        assert report.values_trimmed == 2

        cleaned = load_workbook(output, data_only=False)

        try:
            customers = cleaned["Customers"]
            orders = cleaned["Orders"]

            assert customers.max_row == 3
            assert customers["A2"].value == "Alice"
            assert customers["A3"].value == "Bob"
            assert customers["B3"].value == 0.0
            assert orders["B2"].value == "Completed"

            print("All Excel cleaning tests passed!")
            print(report.to_dict())
        finally:
            cleaned.close()


if __name__ == "__main__":
    test_excel_cleaner()