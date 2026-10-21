from dataclasses import replace
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from app.config import DEFAULT_CONFIG
from app.services.excel_cleaner import clean_workbook


def create_test_excel(path):
    wb = Workbook()
    ws = wb.active

    ws.append(["Name", "Status"])
    ws.append([" Alice ", " Active "])
    ws.append([" Alice ", " Active "])
    ws.append([None, None])
    ws.append(["Bob", "Inactive"])

    wb.save(path)
    wb.close()


def test_disable_duplicate_removal():
    with TemporaryDirectory() as folder:
        source = Path(folder) / "test.xlsx"
        create_test_excel(source)

        config = replace(
            DEFAULT_CONFIG,
            remove_exact_duplicates=False
        )

        report = clean_workbook(source, config=config)

        assert report.duplicate_rows_removed == 0

        wb = load_workbook(report.output_file)
        try:
            assert wb.active.max_row == 4
            assert wb.active["A2"].value == "Alice"
            assert wb.active["A3"].value == "Alice"
        finally:
            wb.close()


def test_disable_blank_row_removal():
    with TemporaryDirectory() as folder:
        source = Path(folder) / "test.xlsx"
        create_test_excel(source)

        config = replace(
            DEFAULT_CONFIG,
            remove_blank_rows=False
        )

        report = clean_workbook(source, config=config)

        assert report.blank_rows_removed == 0

        wb = load_workbook(report.output_file)
        try:
            assert wb.active.max_row == 4
            assert wb.active["A3"].value == "Unknown"
            assert wb.active["B3"].value == "N/A"
        finally:
            wb.close()


def test_disable_whitespace_trimming():
    with TemporaryDirectory() as folder:
        source = Path(folder) / "test.xlsx"
        create_test_excel(source)

        config = replace(
            DEFAULT_CONFIG,
            trim_whitespace=False
        )

        report = clean_workbook(source, config=config)

        assert report.values_trimmed == 0

        wb = load_workbook(report.output_file)
        try:
            assert wb.active["A2"].value == " Alice "
        finally:
            wb.close()